import io
import tempfile
from pathlib import Path

from app.config import Settings
from app.engine import Engine

DEMO = Path(__file__).resolve().parent.parent / "demo_data"


def make_engine(llm=None):
    s = Settings(data_dir=Path(tempfile.mkdtemp()), embedding_backend="hashing")
    return Engine(s, llm=llm)


def load_demo(engine):
    for f in sorted(DEMO.glob("*.txt")):
        engine.ingest(f.name, f.read_bytes())
    return engine


def make_pdf(pages: list[str]) -> bytes:
    """Build a tiny but valid multi-page text PDF (no external libs)."""
    objs, kids = [], []
    n = 3 + 2 * len(pages)
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    for i in range(len(pages)):
        kids.append(f"{4 + 2 * i} 0 R")
    objs.append(f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(pages):
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {5 + 2 * i} 0 R "
                    f"/Resources << /Font << /F1 3 0 R >> >> >>".encode())
        objs.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offs = []
    for i, o in enumerate(objs, 1):
        offs.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + o + b"\nendobj\n")
    x = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for o in offs:
        out.write(f"{o:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF".encode())
    return out.getvalue()


def make_docx(paras: list[tuple[str, str]]) -> bytes:
    import docx

    d = docx.Document()
    for style, text in paras:
        if style == "h":
            d.add_heading(text, level=1)
        else:
            d.add_paragraph(text)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


class FakeLLM:
    def __init__(self, response):
        self.response, self.calls = response, 0

    def investigate(self, question, evidence):
        self.calls += 1
        return self.response(evidence) if callable(self.response) else self.response
