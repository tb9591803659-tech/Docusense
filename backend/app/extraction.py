"""Text extraction for PDF, DOCX and TXT/MD files, page-aware."""
from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
EST_PAGE_CHARS = 3000


class ExtractionError(Exception):
    """Raised with a user-safe message when a file cannot be read."""


@dataclass
class Page:
    number: int
    text: str


@dataclass
class Extraction:
    pages: list[Page]
    page_basis: str  # "native" (real page breaks) or "estimated"


def _paginate(blocks: list[str]) -> list[Page]:
    pages, cur, size = [], [], 0
    for b in blocks:
        if cur and size + len(b) > EST_PAGE_CHARS:
            pages.append(Page(len(pages) + 1, "\n\n".join(cur)))
            cur, size = [], 0
        cur.append(b)
        size += len(b)
    if cur:
        pages.append(Page(len(pages) + 1, "\n\n".join(cur)))
    return pages


def _pdf(data: bytes) -> Extraction:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("This PDF is password-protected.")
        pages = [Page(i + 1, (p.extract_text() or "").strip()) for i, p in enumerate(reader.pages)]
    except ExtractionError:
        raise
    except Exception as e:  # any parser failure on untrusted input
        raise ExtractionError("The PDF is corrupt or could not be read.") from e
    return Extraction(pages, "native")


def _docx(data: bytes) -> Extraction:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:  # BadZipFile, PackageNotFoundError, malformed XML, ...
        raise ExtractionError("The DOCX file is corrupt or could not be read.") from e
    blocks = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "") if p.style is not None else ""
        blocks.append(f"# {t}" if style.lower().startswith(("heading", "title")) else t)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return Extraction(_paginate(blocks), "estimated")


def _text(data: bytes) -> Extraction:
    try:
        s = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        s = data.decode("latin-1")
    if "\f" in s:  # form feed = explicit page break
        parts = [p.strip() for p in s.split("\f")]
        return Extraction([Page(i + 1, p) for i, p in enumerate(parts)], "native")
    blocks = [b.strip() for b in s.replace("\r\n", "\n").split("\n\n") if b.strip()]
    return Extraction(_paginate(blocks), "estimated")


def extract(filename: str, data: bytes) -> Extraction:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ExtractionError(f"Unsupported file type '{ext or 'none'}'. Use PDF, DOCX, TXT or MD.")
    if not data:
        raise ExtractionError("The file is empty.")
    result = {".pdf": _pdf, ".docx": _docx}.get(ext, _text)(data)
    if not any(p.text.strip() for p in result.pages):
        raise ExtractionError("No extractable text found (scanned or image-only document?).")
    return result
